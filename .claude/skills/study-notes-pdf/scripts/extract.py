"""Extract plain text from a source file (pdf, pptx, docx, txt) into a .txt file.

Usage: python3 -I extract.py <source> <out.txt>

PDF pages are separated by "=== PAGE n ===" markers and pptx slides by
"=== SLIDE n ===" so you can cite and navigate the source while reading.
Speaker notes in pptx are included under "[notes]" because lecturers often put
exam-relevant detail there.
"""
import subprocess
import sys
from pathlib import Path


def pdf(src):
    import pypdf
    n = len(pypdf.PdfReader(str(src)).pages)
    parts = []
    for i in range(1, n + 1):
        t = subprocess.run(["pdftotext", "-layout", "-f", str(i), "-l", str(i), str(src), "-"],
                           capture_output=True, text=True).stdout
        parts.append(f"=== PAGE {i} ===\n{t.strip()}")
    return "\n\n".join(parts)


def pptx(src):
    from pptx import Presentation
    parts = []
    for i, slide in enumerate(Presentation(str(src)).slides, 1):
        lines = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append("  " * para.level + "- " + t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines.append("[notes] " + note)
        parts.append(f"=== SLIDE {i} ===\n" + "\n".join(lines))
    return "\n\n".join(parts)


def docx(src):
    import docx as d
    doc = d.Document(str(src))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for tb in doc.tables:
        for row in tb.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def main():
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    ext = src.suffix.lower()
    text = {".pdf": pdf, ".pptx": pptx, ".docx": docx}.get(ext, lambda s: s.read_text(errors="replace"))(src)
    out.write_text(text)
    print(f"{src.name}: {len(text):,} chars -> {out}")
    if ext == ".pdf" and len(text.replace("=== PAGE", "").split()) < 50:
        print("WARNING: almost no text; the PDF is probably scanned and needs OCR.")


if __name__ == "__main__":
    main()
